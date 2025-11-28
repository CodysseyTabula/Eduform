#!/usr/bin/env python3
"""
통합 IEP 생성 파이프라인 실행 스크립트

실제 LLM을 호출하여 데이터를 생성하고, 최종적으로 DOCX 파일까지 생성하는 통합 스크립트입니다.

프로세스:
1. AI 모듈 호출 (실제 LLM)
   - recommend_goal → goals.json 생성
   - recommend_weekly → weekly_content.json 생성
   - recommend_weekly_material → weekly_material_recommendations.json 생성
2. 데이터 형식 변환 및 검증
3. JSON 파일 저장 (storage/iep/{iep_version_id}/)
4. DOCX 파일 생성 및 검증
"""

import sys
import os
import json
import time
import uuid
import subprocess
import shutil
from pathlib import Path
from typing import Dict, Any, List

# backend 디렉토리를 Python 경로에 추가
backend_dir = Path(__file__).parent.parent
sys.path.insert(0, str(backend_dir))
scripts_dir = Path(__file__).parent

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from app.services.file_storage import save_json_file as save_json_file_service


# 도메인 매핑
KOREAN_DOMAIN_MAPPING = {
    "듣기말하기": "listeningSpeaking",
    "읽기": "reading",
    "쓰기": "writing",
    "문법": "grammar",
    "문학": "literature",
    "매체": "mediaLiteracy"
}

MATH_DOMAIN_MAPPING = {
    "수와 연산": "numbersOperations",
    "변화와 관계": "changeAndRelations",
    "도형과 측정": "geometryMeasurement",
    "자료와 가능성": "dataAndProbability"
}

ALL_DOMAIN_MAPPING = {**KOREAN_DOMAIN_MAPPING, **MATH_DOMAIN_MAPPING}


def check_file_exists(filepath: str) -> bool:
    """파일 존재 여부 확인"""
    if Path(scripts_dir / filepath).exists():
        return True
    if Path(backend_dir / filepath).exists():
        return True
    return False


def get_file_path(filepath: str) -> Path:
    """파일 경로 반환"""
    scripts_path = scripts_dir / filepath
    backend_path = backend_dir / filepath
    if scripts_path.exists():
        return scripts_path
    if backend_path.exists():
        return backend_path
    raise FileNotFoundError(f"파일을 찾을 수 없습니다: {filepath}")


def run_ai_module(module_name: str, required_files: List[str]) -> bool:
    """AI 모듈 실행"""
    # 필수 파일 확인
    missing_files = [f for f in required_files if not check_file_exists(f)]
    if missing_files:
        print(f"  ❌ 오류: 다음 파일들이 없습니다: {', '.join(missing_files)}")
        return False
    
    # 모듈 실행
    try:
        result = subprocess.run(
            [sys.executable, "-m", module_name],
            cwd=backend_dir,
            check=True,
            capture_output=False
        )
        return True
    except subprocess.CalledProcessError as e:
        print(f"  ❌ 모듈 실행 실패: {e}")
        return False
    except KeyboardInterrupt:
        print(f"  ⚠️  사용자에 의해 중단됨")
        return False


def load_student_profile(file_path: str = None) -> dict:
    """학생 프로필 로드"""
    if file_path is None:
        scripts_profile = scripts_dir / "student_info.json"
        backend_profile = backend_dir / "student_info.json"
        
        if scripts_profile.exists():
            file_path = scripts_profile
        elif backend_profile.exists():
            file_path = backend_profile
        else:
            raise FileNotFoundError(f"학생 프로필 파일을 찾을 수 없습니다: {scripts_profile} 또는 {backend_profile}")
    
    with open(file_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def convert_goals_to_docx_format(goals_json: Dict) -> Dict[str, Dict[str, str]]:
    """goals.json 형식을 DOCX 생성 형식으로 변환"""
    docx_goals = {}
    
    # goals.json 형식: {"file_type": "goal", "annual_reading_goal": "...", "semester_reading_goal": "..."}
    # DOCX 형식: {"reading": {"annual_goal": "...", "semester_goal": "..."}}
    
    for key, value in goals_json.items():
        if key == "file_type" or not value:
            continue
        
        # 키 파싱: "annual_reading_goal" -> ("reading", "annual_goal")
        if key.startswith("annual_"):
            domain_key = key.replace("annual_", "").replace("_goal", "")
            goal_type = "annual_goal"
        elif key.startswith("semester_"):
            domain_key = key.replace("semester_", "").replace("_goal", "")
            goal_type = "semester_goal"
        else:
            continue
        
        if domain_key not in docx_goals:
            docx_goals[domain_key] = {}
        docx_goals[domain_key][goal_type] = value
    
    return docx_goals


def convert_weekly_content_to_docx_format(weekly_json: Dict) -> Dict[str, List[Dict[str, Any]]]:
    """weekly_content.json 형식을 DOCX 생성 형식으로 변환"""
    docx_weekly = {}
    
    # weekly_content.json 형식: {"file_type": "weekly", "reading_weeklyContent": ["...", "..."]}
    # 또는 LLM 출력형: {"reading": [{"week": 1, "content": "..."}], ...}
    # DOCX 형식: {"reading": [{"week": 1, "content": "..."}, ...]}
    
    for key, value in weekly_json.items():
        if key == "file_type" or not isinstance(value, list):
            continue
        
        if key.endswith("_weeklyContent"):
            domain_key = key.replace("_weeklyContent", "")
            
            # 주차별 데이터로 변환
            weekly_list = []
            for week_idx, content in enumerate(value, 1):
                weekly_list.append({
                    "week": week_idx,
                    "content": content
                })
            
            docx_weekly[domain_key] = weekly_list
        elif key in ALL_DOMAIN_MAPPING.values():
            # 후방 호환: 도메인 키 자체가 있고 [{week, content}] 형식인 경우
            weekly_list = []
            for item in value:
                if isinstance(item, dict):
                    week = item.get("week")
                    content = item.get("content", "")
                    weekly_list.append({
                        "week": week if week is not None else len(weekly_list) + 1,
                        "content": content
                    })
                elif isinstance(item, str):
                    weekly_list.append({
                        "week": len(weekly_list) + 1,
                        "content": item
                    })
            if weekly_list:
                docx_weekly[key] = weekly_list
    
    return docx_weekly


def convert_materials_to_docx_format(materials_json: Dict) -> Dict[str, List[Dict[str, Any]]]:
    """weekly_material_recommendations.json 형식을 DOCX 생성 형식으로 변환"""
    docx_materials = {}
    
    # materials 형식: {"file_type": "material", "reading_weekly_material": [{"title": "...", "content_url": "..."}, ...]}
    # DOCX 형식: {"reading": [{"week": 1, "material_url": "...", "title": "..."}, ...]}
    
    for key, value in materials_json.items():
        if key == "file_type" or not isinstance(value, list):
            continue
        
        if key.endswith("_weekly_material"):
            domain_key = key.replace("_weekly_material", "")
            
            # 주차별 데이터로 변환
            materials_list = []
            for week_idx, material in enumerate(value, 1):
                if isinstance(material, dict):
                    material_url = material.get("content_url", "")
                    material_title = material.get("title", "")
                else:
                    material_url = ""
                    material_title = ""
                
                materials_list.append({
                    "week": week_idx,
                    "material_url": material_url,
                    "title": material_title
                })
            
            docx_materials[domain_key] = materials_list
    
    return docx_materials


def validate_data_format(goals_data: dict, weekly_plan_data: dict, weekly_materials_data: dict) -> bool:
    """데이터 형식 검증"""
    try:
        # Goals 형식 검증
        for domain_key, goal_data in goals_data.items():
            assert isinstance(goal_data, dict), f"{domain_key}: dict가 아님"
            assert "annual_goal" in goal_data, f"{domain_key}: annual_goal 없음"
            assert "semester_goal" in goal_data, f"{domain_key}: semester_goal 없음"
        
        # Weekly Plan 형식 검증
        for domain_key, weekly_list in weekly_plan_data.items():
            assert isinstance(weekly_list, list), f"{domain_key}: list가 아님"
            assert len(weekly_list) == 20, f"{domain_key}: 20주차가 아님 (실제: {len(weekly_list)}주차)"
            for item in weekly_list:
                assert isinstance(item, dict), f"{domain_key}: item이 dict가 아님"
                assert "week" in item, f"{domain_key}: week 필드 없음"
                assert "content" in item, f"{domain_key}: content 필드 없음"
        
        # Weekly Materials 형식 검증
        for domain_key, materials_list in weekly_materials_data.items():
            assert isinstance(materials_list, list), f"{domain_key}: list가 아님"
            assert len(materials_list) == 20, f"{domain_key}: 20주차가 아님 (실제: {len(materials_list)}주차)"
            for item in materials_list:
                assert isinstance(item, dict), f"{domain_key}: item이 dict가 아님"
                assert "week" in item, f"{domain_key}: week 필드 없음"
                assert "material_url" in item, f"{domain_key}: material_url 필드 없음"
        
        return True
    except AssertionError as e:
        print(f"  ✗ 데이터 형식 검증 실패: {e}")
        return False


def save_json_file(iep_version_id: str, file_type: str, content: dict) -> str:
    """
    JSON 파일 저장
    
    파일 경로 패턴: {STORAGE_PATH}/iep/{iep_version_id}/{file_type}-{ts}.json
    """
    # 서비스 함수 사용 (일관성 유지)
    return save_json_file_service(iep_version_id, file_type, content)


def generate_docx(goals_data: dict, weekly_plan_data: dict, weekly_materials_data: dict, 
                  output_path: str = None, iep_version_id: str = None) -> str:
    """DOCX 파일 생성"""
    # 출력 경로 설정
    if output_path is None:
        if iep_version_id:
            output_dir = backend_dir / "storage" / "iep" / iep_version_id
        else:
            output_dir = backend_dir / "test_output"
        output_dir.mkdir(parents=True, exist_ok=True)
        timestamp = int(time.time())
        output_path = str(output_dir / f"iep-{timestamp}.docx")
    
    # 문서 생성
    doc = Document()
    
    # 문서 제목
    title = doc.add_heading('개별화 교육 계획 (IEP)', level=1)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # 도메인 한글 매핑
    domain_names = {
        "listeningSpeaking": "국어 - 듣기말하기",
        "reading": "국어 - 읽기",
        "writing": "국어 - 쓰기",
        "grammar": "국어 - 문법",
        "literature": "국어 - 문학",
        "mediaLiteracy": "국어 - 매체",
        "numbersOperations": "수학 - 수와 연산",
        "changeAndRelations": "수학 - 변화와 관계",
        "geometryMeasurement": "수학 - 도형과 측정",
        "dataAndProbability": "수학 - 자료와 가능성",
    }
    
    # 1. 교육 목표 섹션
    doc.add_heading('1. 교육 목표', level=2)
    for domain_key, domain_name in domain_names.items():
        if domain_key in goals_data:
            domain_goals = goals_data[domain_key]
            doc.add_heading(domain_name, level=3)
            if isinstance(domain_goals, dict) and "annual_goal" in domain_goals:
                doc.add_paragraph(
                    f"📌 연간 목표: {domain_goals['annual_goal']}",
                    style='List Bullet'
                )
            if isinstance(domain_goals, dict) and "semester_goal" in domain_goals:
                doc.add_paragraph(
                    f"🎯 학기 목표: {domain_goals['semester_goal']}",
                    style='List Bullet'
                )
            doc.add_paragraph()
    
    if not any(key in goals_data for key in domain_names.keys()):
        doc.add_paragraph("⚠️ 목표 데이터가 없습니다.", style='Intense Quote')
    
    # 2. 주차별 학습 계획 섹션
    doc.add_heading('2. 주차별 학습 계획 (20주)', level=2)
    for domain_key, domain_name in domain_names.items():
        if domain_key in weekly_plan_data:
            weekly_content = weekly_plan_data[domain_key]
            doc.add_heading(domain_name, level=3)
            if isinstance(weekly_content, list):
                for item in weekly_content:
                    if isinstance(item, dict):
                        week = item.get('week', '?')
                        content = item.get('content', '내용 없음')
                        doc.add_paragraph(
                            f"[{week}주차] {content}",
                            style='List Number'
                        )
            doc.add_paragraph()
    
    if not any(key in weekly_plan_data for key in domain_names.keys()):
        doc.add_paragraph("⚠️ 주차별 학습 계획 데이터가 없습니다.", style='Intense Quote')
    
    # 3. 주차별 학습 자료 섹션
    doc.add_heading('3. 주차별 학습 자료', level=2)
    for domain_key, domain_name in domain_names.items():
        if domain_key in weekly_materials_data:
            materials_list = weekly_materials_data[domain_key]
            doc.add_heading(domain_name, level=3)
            if isinstance(materials_list, list):
                for item in materials_list:
                    if isinstance(item, dict):
                        week = item.get('week', '?')
                        material_url = item.get('material_url', '링크 없음')
                        material_title = item.get('title', '')
                        
                        # 제목이 있으면 제목과 링크를 함께 표시, 없으면 링크만 표시
                        if material_title:
                            doc.add_paragraph(
                                f"[{week}주차] {material_title} - {material_url}",
                                style='List Bullet 2'
                            )
                        else:
                            doc.add_paragraph(
                                f"[{week}주차] {material_url}",
                                style='List Bullet 2'
                            )
            doc.add_paragraph()
    
    if not any(key in weekly_materials_data for key in domain_names.keys()):
        doc.add_paragraph("⚠️ 학습 자료 데이터가 없습니다.", style='Intense Quote')
    
    # 파일 저장
    doc.save(output_path)
    return output_path


def main():
    """전체 파이프라인 실행"""
    print("=" * 60)
    print("통합 IEP 생성 파이프라인 실행")
    print("=" * 60)
    print()
    
    # 환경 변수 확인
    from dotenv import load_dotenv
    load_dotenv()
    
    openai_key = os.getenv("OPENAI_API_KEY")
    pinecone_key = os.getenv("PINECONE_API_KEY")
    
    if not openai_key:
        print("❌ 오류: OPENAI_API_KEY가 설정되지 않았습니다.")
        print("   .env 파일에 OPENAI_API_KEY를 추가하세요.")
        return False
    
    # IEP 버전 ID 생성
    iep_version_id = str(uuid.uuid4())
    print(f"IEP 버전 ID: {iep_version_id}")
    print()
    
    try:
        # 1. 학생 프로필 로드
        print("[1단계] 학생 프로필 로드...")
        student_profile = load_student_profile()
        print(f"  ✓ 학생 프로필 로드 완료: {student_profile.get('name', 'N/A')}")
        print()
        
        # 2. AI 모듈 호출 - 목표 생성
        print("[2단계] AI 모듈 데이터 생성...")
        print("  - Goals 생성 중...")
        if not run_ai_module("ai_module.recommend_goal", ["student_info.json"]):
            return False
        
        # goals.json 읽기 및 변환
        goals_file = get_file_path("goals.json")
        with open(goals_file, 'r', encoding='utf-8') as f:
            goals_json = json.load(f)
        goals_data = convert_goals_to_docx_format(goals_json)
        print(f"    ✓ Goals 생성 완료: {len(goals_data)}개 도메인")
        print()
        
        # 3. AI 모듈 호출 - 주차별 학습 내용 생성
        print("  - Weekly Plan 생성 중...")
        if not run_ai_module("ai_module.recommend_weekly", ["student_info.json", "goals.json"]):
            return False
        
        # weekly_content.json 읽기 및 변환
        weekly_file = get_file_path("weekly_content.json")
        with open(weekly_file, 'r', encoding='utf-8') as f:
            weekly_json = json.load(f)
        weekly_plan_data = convert_weekly_content_to_docx_format(weekly_json)
        print(f"    ✓ Weekly Plan 생성 완료: {len(weekly_plan_data)}개 도메인")
        print()
        
        # 4. AI 모듈 호출 - 주차별 학습 자료 추천 (선택적)
        weekly_materials_data = {}
        if pinecone_key:
            print("  - Weekly Materials 생성 중...")
            if run_ai_module("ai_module.recommend_weekly_material", ["weekly_content.json"]):
                # weekly_material_recommendations.json 읽기 및 변환
                materials_file = get_file_path("weekly_material_recommendations.json")
                with open(materials_file, 'r', encoding='utf-8') as f:
                    materials_json = json.load(f)
                weekly_materials_data = convert_materials_to_docx_format(materials_json)
                print(f"    ✓ Weekly Materials 생성 완료: {len(weekly_materials_data)}개 도메인")
            else:
                print("    ⚠️  Weekly Materials 생성 실패 (계속 진행)")
        else:
            print("  - Weekly Materials 건너뜀 (PINECONE_API_KEY 없음)")
        print()
        
        # 5. 데이터 형식 검증
        print("[3단계] 데이터 형식 검증...")
        if not validate_data_format(goals_data, weekly_plan_data, weekly_materials_data):
            print("  ✗ 데이터 형식 검증 실패")
            return False
        print("  ✓ 모든 데이터 형식 검증 통과")
        print()
        
        # 6. JSON 파일 저장
        print("[4단계] JSON 파일 저장...")
        goals_path = save_json_file(iep_version_id, "goals", goals_data)
        print(f"  ✓ Goals 저장: {goals_path}")
        
        weekly_plan_path = save_json_file(iep_version_id, "weekly_plan", weekly_plan_data)
        print(f"  ✓ Weekly Plan 저장: {weekly_plan_path}")
        
        if weekly_materials_data:
            weekly_materials_path = save_json_file(iep_version_id, "weekly_materials", weekly_materials_data)
            print(f"  ✓ Weekly Materials 저장: {weekly_materials_path}")
        print()
        
        # 7. DOCX 파일 생성
        print("[5단계] DOCX 파일 생성...")
        docx_path = generate_docx(
            goals_data,
            weekly_plan_data,
            weekly_materials_data,
            iep_version_id=iep_version_id
        )
        
        # 파일 검증
        docx_file = Path(docx_path)
        assert docx_file.exists(), "DOCX 파일이 생성되지 않음"
        assert docx_file.stat().st_size > 0, "DOCX 파일이 비어있음"
        
        # 문서 내용 확인
        loaded_doc = Document(docx_path)
        paragraphs = [p.text for p in loaded_doc.paragraphs]
        full_text = "\n".join(paragraphs)
        
        assert "개별화 교육 계획" in full_text or "IEP" in full_text, "제목 없음"
        assert "목표" in full_text, "목표 섹션 없음"
        assert "학습 계획" in full_text or "주차" in full_text, "학습 계획 섹션 없음"
        
        print(f"  ✓ DOCX 파일 생성 완료: {docx_path}")
        print(f"  ✓ 파일 크기: {docx_file.stat().st_size:,} bytes")
        print(f"  ✓ 문서 단락 수: {len(loaded_doc.paragraphs)}개")
        
        # test_output에도 복사
        test_output_dir = backend_dir / "test_output"
        test_output_dir.mkdir(parents=True, exist_ok=True)
        test_output_path = test_output_dir / docx_file.name
        shutil.copy2(docx_path, test_output_path)
        print(f"  ✓ test_output에도 저장 완료: {test_output_path}")
        print()
        
        # 완료 메시지
        print("=" * 60)
        print("✅ 전체 파이프라인 완료!")
        print("=" * 60)
        print()
        print("생성된 파일:")
        print(f"  ✓ {goals_path}")
        print(f"  ✓ {weekly_plan_path}")
        if weekly_materials_data:
            print(f"  ✓ {weekly_materials_path}")
        print(f"  ✓ {docx_path}")
        print(f"  ✓ {test_output_path} (test_output 복사본)")
        print()
        
        return True
        
    except FileNotFoundError as e:
        print(f"❌ 오류: {e}")
        print("\n해결 방법:")
        print("  - student_info.json 파일이 scripts/ 또는 backend/ 폴더에 있는지 확인하세요")
        print("  - 또는 --profile 옵션으로 파일 경로를 지정하세요")
        return False
        
    except Exception as e:
        print(f"❌ 오류 발생: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(description="통합 IEP 생성 파이프라인 실행")
    parser.add_argument(
        "--profile",
        type=str,
        help="학생 프로필 JSON 파일 경로 (기본값: scripts/student_info.json 또는 backend/student_info.json)"
    )
    parser.add_argument(
        "--output",
        type=str,
        help="DOCX 출력 파일 경로 (기본값: 자동 생성)"
    )
    
    args = parser.parse_args()
    
    # 프로필 파일 경로 설정
    if args.profile:
        original_load = load_student_profile
        def load_student_profile_wrapper():
            return original_load(args.profile)
        load_student_profile = load_student_profile_wrapper
    
    # 출력 경로 설정
    if args.output:
        original_generate = generate_docx
        def generate_docx_wrapper(goals, plan, materials, output_path=None, iep_version_id=None):
            return original_generate(goals, plan, materials, args.output, iep_version_id)
        generate_docx = generate_docx_wrapper
    
    try:
        success = main()
        sys.exit(0 if success else 1)
    except KeyboardInterrupt:
        print("\n\n⚠️  사용자에 의해 중단되었습니다.")
        sys.exit(1)
