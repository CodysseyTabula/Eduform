"""
DOCX 문서 생성 및 다운로드 테스트

IEP 버전의 DOCX 파일 생성 및 다운로드 API 테스트
"""
import io
import json
import pytest
from docx import Document


def _create_all_iep_files(client, iep_version_id, student_profile):
    """DOCX 생성 전에 3개 file_type 모두 준비"""
    profile_json = json.dumps(student_profile, ensure_ascii=False).encode('utf-8')
    for file_type in ["goals", "weekly_plan", "weekly_materials"]:
        files = {
            "file": ("student_profile.json", profile_json, "application/json")
        }
        data = {
            "iep_version_id": str(iep_version_id),
            "file_type": file_type
        }
        response = client.post("/iep-files", files=files, data=data)
        assert response.status_code == 201


@pytest.mark.integration
def test_download_iep_docx(
    client,
    sample_iep_version,
    sample_student_profile,
    mock_ai_generators
):
    """
    GET /iep-versions/{iep_version_id}/docx - DOCX 다운로드 테스트
    
    검증:
    - 사전 조건: 3개 IEP 파일 존재
    - 200 OK 응답
    - Content-Type이 DOCX 형식
    - 다운로드된 파일이 유효한 DOCX
    """
    _create_all_iep_files(
        client,
        sample_iep_version.id,
        sample_student_profile
    )
    
    # 2. DOCX 다운로드
    response = client.get(f"/iep-versions/{sample_iep_version.id}/docx")
    
    # 응답 확인
    assert response.status_code == 200
    
    # Content-Type 확인
    content_type = response.headers.get("content-type")
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in content_type
    
    # Content-Disposition 확인
    content_disposition = response.headers.get("content-disposition")
    assert "attachment" in content_disposition
    assert f"IEP_{sample_iep_version.id}.docx" in content_disposition
    
    # DOCX 파일 유효성 검증
    docx_bytes = response.content
    assert len(docx_bytes) > 0
    
    # python-docx로 파싱 가능한지 확인
    try:
        doc = Document(io.BytesIO(docx_bytes))
        # 문서가 최소한의 내용을 포함하는지 확인
        assert len(doc.paragraphs) > 0
    except Exception as e:
        pytest.fail(f"Failed to parse DOCX: {e}")


@pytest.mark.integration
def test_docx_content_structure(
    client,
    sample_iep_version,
    sample_student_profile,
    mock_ai_generators
):
    """
    DOCX 문서 내용 구조 검증
    
    검증:
    - 문서에 제목이 포함됨
    - IEP 메타데이터가 포함됨
    - 목표/학습계획/자료 섹션이 포함됨
    """
    _create_all_iep_files(
        client,
        sample_iep_version.id,
        sample_student_profile
    )
    
    # 2. DOCX 다운로드
    response = client.get(f"/iep-versions/{sample_iep_version.id}/docx")
    assert response.status_code == 200
    
    # 3. DOCX 내용 확인
    doc = Document(io.BytesIO(response.content))
    
    # 전체 텍스트 추출
    full_text = "\n".join([p.text for p in doc.paragraphs])
    
    # 제목 확인
    assert "개별화 교육 계획" in full_text or "IEP" in full_text
    
    # 메타데이터 확인
    assert "2024" in full_text  # year
    assert "1학기" in full_text  # semester
    
    # 섹션 확인 (목표/계획/자료 중 하나라도)
    has_content = (
        "목표" in full_text or
        "학습" in full_text or
        "자료" in full_text or
        "주차" in full_text
    )
    assert has_content


@pytest.mark.integration
def test_docx_download_without_files(client, sample_iep_version):
    """
    IEP 파일 없이 DOCX 다운로드 시도 (실패 케이스)
    
    검증:
    - 404 또는 500 에러 반환
    - 적절한 에러 메시지
    """
    # IEP 파일을 생성하지 않고 바로 DOCX 다운로드 시도
    response = client.get(f"/iep-versions/{sample_iep_version.id}/docx")
    
    # 파일이 없으므로 에러 응답 예상
    assert response.status_code in [404, 500]
    
    # JSON 에러 메시지 확인
    data = response.json()
    assert "detail" in data


@pytest.mark.integration
def test_docx_download_nonexistent_version(client):
    """
    존재하지 않는 IEP 버전의 DOCX 다운로드 시도 (실패 케이스)
    
    검증:
    - 404 에러 반환
    """
    import uuid
    fake_id = str(uuid.uuid4())
    
    response = client.get(f"/iep-versions/{fake_id}/docx")
    
    assert response.status_code == 404
    data = response.json()
    assert "detail" in data
