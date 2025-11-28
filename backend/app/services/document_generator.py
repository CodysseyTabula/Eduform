"""
IEP DOCX 문서 생성 서비스

IEP JSON 파일(student_info, goal, weekly_content, weekly_material)을 읽어
DOCX 문서를 생성하고 다운로드를 제공합니다.
"""

from __future__ import annotations

import os
import time
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.iep_file import IEPFile
from app.models.iep_version import IEPVersion
from app.services.file_storage import load_json_file, FileStorageError


# ==================== 예외 정의 ====================

class DocumentGenerationError(Exception):
    """문서 생성 관련 기본 예외"""
    pass


class DocumentNotFoundError(DocumentGenerationError):
    """필요한 IEP 파일을 찾을 수 없을 때"""
    pass


class InvalidDocumentDataError(DocumentGenerationError):
    """JSON 데이터가 유효하지 않을 때"""
    pass


# ==================== 핵심 함수 ====================

def generate_docx_for_iep(
    db: Session,
    iep_version_id: UUID,
    output_path: str | None = None
) -> str:
    """
    IEP DOCX 파일 생성
    
    Args:
        db: 데이터베이스 세션
        iep_version_id: IEP 버전 ID (UUID)
        output_path: 저장할 파일 경로 (없으면 자동 생성)
    
    Returns:
        str: 생성된 DOCX 파일의 절대 경로
    
    Raises:
        DocumentNotFoundError: IEP 버전이나 필수 파일이 없을 때
        InvalidDocumentDataError: JSON 데이터가 유효하지 않을 때
        DocumentGenerationError: 기타 문서 생성 오류
    """
    try:
        # 1. IEP 버전 조회
        iep_version = db.query(IEPVersion).filter(
            IEPVersion.id == iep_version_id
        ).first()
        
        if not iep_version:
            raise DocumentNotFoundError(
                f"IEP version not found: {iep_version_id}"
            )
        
        # 2. IEP 파일 메타데이터 조회 (4개 파일 필요)
        iep_files = db.query(IEPFile).filter(
            IEPFile.iep_version_id == iep_version_id
        ).all()
        
        file_map = {f.file_type: f.file_path for f in iep_files}
        
        # 필수 파일 확인
        required_types = ["student_info", "goal", "weekly_content", "weekly_material"]
        missing_types = [ft for ft in required_types if ft not in file_map]
        
        if missing_types:
            raise DocumentNotFoundError(
                f"Missing required IEP files: {', '.join(missing_types)}"
            )
        
        # 3. JSON 파일 로드
        try:
            student_info_data = load_json_file(file_map["student_info"])
            goals_data = load_json_file(file_map["goal"])
            weekly_plan_data = load_json_file(file_map["weekly_content"])
            weekly_materials_data = load_json_file(file_map["weekly_material"])
        except FileStorageError as e:
            raise DocumentNotFoundError(f"Failed to load IEP JSON files: {e}") from e
        
        # 4. DOCX 문서 생성
        doc = Document()
        
        # 문서 제목
        title = doc.add_heading('개별화 교육 계획 (IEP)', level=1)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # IEP 메타데이터 섹션
        _add_metadata_section(doc, iep_version, student_info_data)
        
        # 연간/학기 목표 섹션 (선택된 도메인만 출력)
        _add_goals_section(doc, goals_data, student_info_data)
        
        # 주차별 학습 계획 섹션 (선택된 도메인만 출력)
        _add_weekly_plan_section(doc, weekly_plan_data, student_info_data)
        
        # 주차별 학습 자료 섹션 (선택된 도메인만 출력)
        _add_weekly_materials_section(doc, weekly_materials_data, student_info_data)
        
        # 5. 파일 저장
        if not output_path:
            # 자동 경로 생성: {STORAGE_PATH}/iep/{iep_version_id}/iep-{ts}.docx
            version_dir = Path(settings.storage_path) / "iep" / str(iep_version_id)
            version_dir.mkdir(parents=True, exist_ok=True)
            timestamp = int(time.time())
            output_path = str(version_dir / f"iep-{timestamp}.docx")
        
        doc.save(output_path)
        
        return str(Path(output_path).absolute())
    
    except (DocumentNotFoundError, InvalidDocumentDataError):
        raise
    except Exception as e:
        raise DocumentGenerationError(f"Failed to generate DOCX: {e}") from e


def generate_docx_stream(db: Session, iep_version_id: UUID) -> BytesIO:
    """
    IEP DOCX 파일을 메모리 스트림으로 생성 (다운로드용)
    
    Args:
        db: 데이터베이스 세션
        iep_version_id: IEP 버전 ID (UUID)
    
    Returns:
        BytesIO: DOCX 파일 바이너리 스트림
    
    Raises:
        DocumentNotFoundError, InvalidDocumentDataError, DocumentGenerationError
    """
    try:
        # 1. IEP 버전 조회
        iep_version = db.query(IEPVersion).filter(
            IEPVersion.id == iep_version_id
        ).first()
        
        if not iep_version:
            raise DocumentNotFoundError(
                f"IEP version not found: {iep_version_id}"
            )
        
        # 2. IEP 파일 메타데이터 조회
        iep_files = db.query(IEPFile).filter(
            IEPFile.iep_version_id == iep_version_id
        ).all()
        
        file_map = {f.file_type: f.file_path for f in iep_files}
        
        # 필수 파일 확인
        required_types = ["student_info", "goal", "weekly_content", "weekly_material"]
        missing_types = [ft for ft in required_types if ft not in file_map]
        
        if missing_types:
            raise DocumentNotFoundError(
                f"Missing required IEP files: {', '.join(missing_types)}"
            )
        
        # 3. JSON 파일 로드
        try:
            student_info_data = load_json_file(file_map["student_info"])
            goals_data = load_json_file(file_map["goal"])
            weekly_plan_data = load_json_file(file_map["weekly_content"])
            weekly_materials_data = load_json_file(file_map["weekly_material"])
        except FileStorageError as e:
            raise DocumentNotFoundError(f"Failed to load IEP JSON files: {e}") from e
        
        # 4. DOCX 문서 생성
        doc = Document()
        
        # 문서 제목
        title = doc.add_heading('개별화 교육 계획 (IEP)', level=1)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # IEP 메타데이터 섹션
        _add_metadata_section(doc, iep_version, student_info_data)
        
        # 연간/학기 목표 섹션 (선택된 도메인만 출력)
        _add_goals_section(doc, goals_data, student_info_data)
        
        # 주차별 학습 계획 섹션 (선택된 도메인만 출력)
        _add_weekly_plan_section(doc, weekly_plan_data, student_info_data)
        
        # 주차별 학습 자료 섹션 (선택된 도메인만 출력)
        _add_weekly_materials_section(doc, weekly_materials_data, student_info_data)
        
        # 5. 메모리 스트림에 저장
        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        
        return stream
    
    except (DocumentNotFoundError, InvalidDocumentDataError):
        raise
    except Exception as e:
        raise DocumentGenerationError(f"Failed to generate DOCX stream: {e}") from e


# ==================== 내부 헬퍼 함수 ====================

def _add_metadata_section(doc: Document, iep_version: IEPVersion, student_info: dict[str, Any] | None = None) -> None:
    """IEP 메타데이터 섹션 추가"""
    doc.add_heading('IEP 정보', level=2)
    
    # 메타데이터를 단락으로 출력 (테스트에서 paragraphs로 추출 가능하도록)
    doc.add_paragraph(f"학년도: {iep_version.year}", style='List Bullet')
    doc.add_paragraph(f"학기: {iep_version.semester}", style='List Bullet')
    doc.add_paragraph(f"학년: {iep_version.grade}", style='List Bullet')
    doc.add_paragraph(f"생성일: {iep_version.created_at.strftime('%Y-%m-%d %H:%M:%S')}", style='List Bullet')
    
    # 학생 프로필 정보
    if student_info:
        doc.add_paragraph(f"학생 이름: {student_info.get('name', '')}", style='List Bullet')
        doc.add_paragraph(f"생년월일: {student_info.get('birth', '')}", style='List Bullet')
        doc.add_paragraph(f"현재 학기: {student_info.get('current_semester', '')}", style='List Bullet')
        doc.add_paragraph(f"IEP 기간: {student_info.get('start_date', '')} ~ {student_info.get('end_date', '')}", style='List Bullet')
    
    doc.add_paragraph()  # 간격


def _add_goals_section(doc: Document, goals_data: dict[str, Any], student_info_data: dict[str, Any] | None = None) -> None:
    """연간/학기 목표 섹션 추가"""
    doc.add_heading('1. 교육 목표', level=2)
    
    # 도메인 한글 매핑
    domain_names = {
        # 국어 도메인
        "listeningSpeaking": "국어 - 듣기말하기",
        "reading": "국어 - 읽기",
        "writing": "국어 - 쓰기",
        "grammar": "국어 - 문법",
        "literature": "국어 - 문학",
        "mediaLiteracy": "국어 - 매체",
        # 수학 도메인
        "numbersOperations": "수학 - 수와 연산",
        "changeAndRelations": "수학 - 변화와 관계",
        "geometryMeasurement": "수학 - 도형과 측정",
        "dataAndProbability": "수학 - 자료와 가능성",
    }
    
    # 선택된 도메인 추출
    selected_domains = set()
    if student_info_data:
        korean_domains = student_info_data.get("korean_domain", [])
        math_domains = student_info_data.get("math_domain", [])
        # 리스트인 경우 그대로 사용, 문자열인 경우 분리
        if isinstance(korean_domains, str):
            korean_domains = [d.strip() for d in korean_domains.split(",") if d.strip()]
        if isinstance(math_domains, str):
            math_domains = [d.strip() for d in math_domains.split(",") if d.strip()]
        selected_domains.update(korean_domains)
        selected_domains.update(math_domains)
    
    has_any_goals = False
    
    # 선택된 도메인만 출력
    # 실제 저장 형식: {annual_{domainKey}_goal: "...", semester_{domainKey}_goal: "..."}
    for domain_key, domain_name in domain_names.items():
        # 선택된 도메인만 처리 (선택된 도메인이 없으면 모든 도메인 처리 - 후방 호환)
        if selected_domains and domain_key not in selected_domains:
            continue
            
        annual_key = f"annual_{domain_key}_goal"
        semester_key = f"semester_{domain_key}_goal"
        
        annual_goal = goals_data.get(annual_key, '')
        semester_goal = goals_data.get(semester_key, '')
        
        # 목표가 하나라도 있으면 출력
        if annual_goal or semester_goal:
            has_any_goals = True
            
            # 도메인별 소제목
            doc.add_heading(domain_name, level=3)
            
            # 연간 목표
            if annual_goal:
                doc.add_paragraph(
                    f"📌 연간 목표: {annual_goal}",
                    style='List Bullet'
                )
            
            # 학기 목표
            if semester_goal:
                doc.add_paragraph(
                    f"🎯 학기 목표: {semester_goal}",
                    style='List Bullet'
                )
            
            doc.add_paragraph()  # 간격
    
    if not has_any_goals:
        doc.add_paragraph("⚠️ 목표 데이터가 없습니다.", style='Intense Quote')


def _add_weekly_plan_section(doc: Document, weekly_plan_data: dict[str, Any], student_info_data: dict[str, Any] | None = None) -> None:
    """주차별 학습 계획 섹션 추가"""
    doc.add_heading('2. 주차별 학습 계획 (20주)', level=2)
    
    # 도메인 한글 매핑
    domain_names = {
        "listeningSpeaking": "국어 - 듣기말하기",
        "reading": "국어 - 읽기",
        "writing": "국어 - 쓰기",
        "grammar": "국어 - 문법",
        "literature": "국어 - 문학",
        "mediaLiteracy": "국어 - 매체",
        "numbersOperations": "수학 - 수와 연산",
        "changeAndRelations": "수학 - 변화와 관계",
        "geometryMeasurement": "수학 - 도형과 측정",
        "dataAndProbability": "수학 - 자료와 가능성",
    }
    
    # 선택된 도메인 추출
    selected_domains = set()
    if student_info_data:
        korean_domains = student_info_data.get("korean_domain", [])
        math_domains = student_info_data.get("math_domain", [])
        # 리스트인 경우 그대로 사용, 문자열인 경우 분리
        if isinstance(korean_domains, str):
            korean_domains = [d.strip() for d in korean_domains.split(",") if d.strip()]
        if isinstance(math_domains, str):
            math_domains = [d.strip() for d in math_domains.split(",") if d.strip()]
        selected_domains.update(korean_domains)
        selected_domains.update(math_domains)
    
    has_any_content = False
    
    # 선택된 도메인만 출력
    # 실제 저장 형식: {domainKey}_weeklyContent: string[] 또는 {domainKey}: [{week, content}]
    for domain_key, domain_name in domain_names.items():
        # 선택된 도메인만 처리 (선택된 도메인이 없으면 모든 도메인 처리 - 후방 호환)
        if selected_domains and domain_key not in selected_domains:
            continue
        weekly_key = f"{domain_key}_weeklyContent"
        weekly_content = weekly_plan_data.get(weekly_key)
        
        # 후방 호환: 도메인 키로 리스트가 있고 dict 항목을 가진 경우 content만 추출
        if not weekly_content and domain_key in weekly_plan_data:
            raw_list = weekly_plan_data.get(domain_key)
            if isinstance(raw_list, list):
                weekly_content = []
                for item in raw_list:
                    if isinstance(item, dict):
                        weekly_content.append(item.get("content", ""))
                    elif isinstance(item, str):
                        weekly_content.append(item)
        
        # 주차별 내용 출력
        if weekly_content and isinstance(weekly_content, list) and len(weekly_content) > 0:
            # 빈 문자열만 있는 경우 제외
            non_empty_content = [c for c in weekly_content if c and str(c).strip()]
            if non_empty_content:
                has_any_content = True
                
                # 도메인별 소제목
                doc.add_heading(domain_name, level=3)
                
                # 주차별 내용 출력 (1주차부터)
                for week_idx, content in enumerate(weekly_content, start=1):
                    if content and str(content).strip():
                        doc.add_paragraph(
                            f"[{week_idx}주차] {content}",
                            style='List Number'
                        )
                
                doc.add_paragraph()  # 간격
    
    if not has_any_content:
        doc.add_paragraph("⚠️ 주차별 학습 계획 데이터가 없습니다.", style='Intense Quote')


def _add_weekly_materials_section(doc: Document, weekly_materials_data: dict[str, Any], student_info_data: dict[str, Any] | None = None) -> None:
    """주차별 학습 자료 섹션 추가"""
    doc.add_heading('3. 주차별 학습 자료', level=2)
    
    # 도메인 한글 매핑
    domain_names = {
        "listeningSpeaking": "국어 - 듣기말하기",
        "reading": "국어 - 읽기",
        "writing": "국어 - 쓰기",
        "grammar": "국어 - 문법",
        "literature": "국어 - 문학",
        "mediaLiteracy": "국어 - 매체",
        "numbersOperations": "수학 - 수와 연산",
        "changeAndRelations": "수학 - 변화와 관계",
        "geometryMeasurement": "수학 - 도형과 측정",
        "dataAndProbability": "수학 - 자료와 가능성",
    }
    
    # 선택된 도메인 추출
    selected_domains = set()
    if student_info_data:
        korean_domains = student_info_data.get("korean_domain", [])
        math_domains = student_info_data.get("math_domain", [])
        # 리스트인 경우 그대로 사용, 문자열인 경우 분리
        if isinstance(korean_domains, str):
            korean_domains = [d.strip() for d in korean_domains.split(",") if d.strip()]
        if isinstance(math_domains, str):
            math_domains = [d.strip() for d in math_domains.split(",") if d.strip()]
        selected_domains.update(korean_domains)
        selected_domains.update(math_domains)
    
    has_any_materials = False
    
    # 선택된 도메인만 출력
    # 실제 저장 형식: {domainKey}_weekly_material: [{title, reason, content_url, thumbnail_url}, ...]
    # 또는 스펙 형식: [{week: number, materials: [{title, url, keywords, file_type}]}, ...]
    for domain_key, domain_name in domain_names.items():
        # 선택된 도메인만 처리 (선택된 도메인이 없으면 모든 도메인 처리 - 후방 호환)
        if selected_domains and domain_key not in selected_domains:
            continue
        material_key = f"{domain_key}_weekly_material"
        materials_list = weekly_materials_data.get(material_key)
        
        # 주차별 자료 출력
        if materials_list and isinstance(materials_list, list) and len(materials_list) > 0:
            # 빈 배열이 아닌지 확인
            has_content = False
            output_items = []
            
            for idx, item in enumerate(materials_list, start=1):
                if isinstance(item, dict):
                    # 스펙 형식: {week: number, materials: [{title, url, ...}]}
                    if 'materials' in item and isinstance(item['materials'], list) and len(item['materials']) > 0:
                        week = item.get('week', idx)  # week가 없으면 인덱스 사용
                        material = item['materials'][0]  # 첫 번째 자료 사용
                        url = material.get('url', material.get('content_url', '링크 없음'))
                        title = material.get('title', '').strip()
                        if url and url.strip():
                            has_content = True
                            if title:
                                output_items.append((week, f"{title} - {url}"))
                            else:
                                # title이 없으면 URL만 출력
                                output_items.append((week, url))
                    # 예전 형식: {title, reason, content_url, thumbnail_url}
                    elif 'content_url' in item or 'material_url' in item:
                        url = item.get('content_url') or item.get('material_url', '링크 없음')
                        title = item.get('title', '').strip()
                        week = item.get('week', idx)  # week가 없으면 인덱스 사용
                        if url and url.strip():
                            has_content = True
                            if title:
                                output_items.append((week, f"{title} - {url}"))
                            else:
                                # title이 없으면 URL만 출력
                                output_items.append((week, url))
                    # 단순 URL 문자열 배열 형식
                    elif 'url' in item:
                        url = item.get('url', '')
                        week = item.get('week', idx)  # week가 없으면 인덱스 사용
                        if url and url.strip():
                            has_content = True
                            output_items.append((week, url))
                # 문자열인 경우 (직접 URL)
                elif isinstance(item, str) and item.strip():
                    has_content = True
                    output_items.append((idx, item))
            
            if has_content:
                has_any_materials = True
                
                # 도메인별 소제목
                doc.add_heading(domain_name, level=3)
                
                # 주차별 자료 출력
                for week, material_info in output_items:
                    doc.add_paragraph(
                        f"[{week}주차] {material_info}",
                        style='List Bullet 2'
                    )
                
                doc.add_paragraph()  # 간격
    
    if not has_any_materials:
        doc.add_paragraph("⚠️ 학습 자료 데이터가 없습니다.", style='Intense Quote')


__all__ = [
    "generate_docx_for_iep",
    "generate_docx_stream",
    "DocumentGenerationError",
    "DocumentNotFoundError",
    "InvalidDocumentDataError",
]
